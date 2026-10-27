"""FICHE 14 — Sortie : test de recette, point par point."""

import csv
import hashlib
import io
import json
import subprocess
import threading
import xml.etree.ElementTree as ET
from http.server import BaseHTTPRequestHandler, HTTPServer
from pathlib import Path
from typing import ClassVar

import pytest
from conftest import copier_config
from fastapi.testclient import TestClient

from pilote.app import creer_app
from pilote.base import Candidat, Recherche, creer_fabrique
from pilote.config import RACINE
from pilote.config import charger as charger_cfg

NOMS = [
    # nom, ville, pays, origine, sources, coordonnées (lat, lon) ou None, relations, fonction
    ("Fenix Toulouse Handball", "Toulouse", "FR", "registre", ["recherche-entreprises"], (43.6045, 1.4442), [{"type": "maison_mere", "vers": "Groupe Fenix"}], "siège"),
    ("Montpellier Handball", "Montpellier", "FR", "registre", ["recherche-entreprises", "wikidata"], (43.6108, 3.8767), [], ""),
    ("Handball Club Grenoble", "Grenoble", "FR", "presse", ["rss_liberation"], (45.1885, 5.7245), [{"type": "meme_adresse", "vers": "Salle des Alpes"}], ""),
    ("Collectif Handball Lille", "Lille", "FR", "web", ["searxng"], None, [], ""),
    ("=HYPERLINK(\"http://x.example\",\"cliquez\") Club", "Nantes", "FR", "web", ["searxng"], None, [], ""),
    ("Handball Shop Nantes", "Nantes", "FR", "web", ["searxng"], None, [], ""),
]


@pytest.fixture
def cfg(tmp_path):
    return charger_cfg(copier_config(tmp_path))


def peupler(cfg, rid="r14"):
    fabrique = creer_fabrique(cfg.url_base)
    with fabrique.begin() as s:
        s.add(Recherche(id=rid, qui="sednal", sujet={"quoi": "clubs de handball", "pays": "France", "inclure": ["pro"]}, parametres={"verification": True}, etat="terminee"))
        s.flush()
        for nom, ville, pays, origine, sources, coord, rel, fonction in NOMS:
            faits = {"fonction": fonction} if fonction else {}
            if coord:
                faits["coordonnees"] = {"lat": coord[0], "lon": coord[1]}
            s.add(Candidat(recherche_id=rid, nom=nom, genre_source=origine, trouve_par=sources, nb_endroits=len(sources), genre_de_chose="club", ville=ville, pays=pays,
                           site="" if "web" == origine else "club.example", relations=rel, faits=faits))
    return fabrique


def decider(c, rid="r14"):
    ids = {x["nom"]: x["id"] for x in c.get(f"/api/recherches/{rid}/tri").json()["lignes"]}
    c.post(f"/api/recherches/{rid}/decisions", json={"valeur": "garder", "ids": [ids[n[0]] for n in NOMS[:3]] + [ids[NOMS[4][0]]]})
    c.post(f"/api/recherches/{rid}/decisions", json={"valeur": "rejeter", "ids": [ids[NOMS[5][0]]], "motif": "doublon"})
    c.post(f"/api/recherches/{rid}/decisions", json={"valeur": "plus_tard", "ids": [ids[NOMS[3][0]]]})
    return ids


@pytest.fixture
def monde(cfg):
    peupler(cfg)
    with TestClient(creer_app(cfg)) as c:
        decider(c)
        yield c, cfg


def exporter(c, formats, **plus):
    r = c.post("/api/recherches/r14/export", json={"formats": formats, **plus})
    assert r.status_code == 200, r.text
    return r.json()


def dossier(res):
    return Path(res["dossier"])


def verifier_trace(trace, attendu_nombre):
    assert trace["sujet"].startswith("clubs de handball") and trace["recherche"] == "r14" and trace["nombre"] == attendu_nombre
    assert trace["date"][:4].isdigit() and {s["id"] for s in trace["sources"]} >= {"recherche-entreprises", "searxng"}
    assert "gardés" in trace["filtre"]


# --- [ ] Les sept formats de la première version produisent un fichier valide.
def test_les_sept_formats_produisent_un_fichier_valide(monde):
    c, _ = monde
    attendus = {x["id"] for x in c.get("/api/formats").json()}
    assert {"csv", "json", "texte", "geojson", "gexf", "xlsx", "webhook"} <= attendus
    srv = Serveur()
    try:
        res = exporter(c, [{"format": f} for f in ("csv", "json", "texte", "geojson", "gexf", "xlsx")] + [{"format": "webhook", "options": {"url": srv.url}}])
    finally:
        srv.arreter()
    d = dossier(res)
    assert res["erreurs"] == []
    # CSV : trois contenus possibles, deux par défaut
    assert next(csv.reader((d / "resultat.csv").read_text(encoding="utf-8").splitlines())) == ["NOM"]
    lignes = list(csv.reader((d / "resultat_tsites.csv").read_text(encoding="utf-8").splitlines(), delimiter=";"))
    assert lignes[0] == ["nom", "ville", "pays"] and len(lignes) == 5  # quatre noms gardés
    # JSON
    j = json.loads((d / "resultat.json").read_text(encoding="utf-8"))
    assert {"trace", "noms"} == set(j) and len(j["noms"]) == 4 and {"nom", "ville", "origine", "sources", "relations"} <= set(j["noms"][0])
    # Texte
    t = (d / "resultat.txt").read_text(encoding="utf-8").splitlines()
    assert t[0].startswith("# Exporté le") and "Fenix Toulouse Handball" in t
    # GeoJSON, GEXF, Excel : voir leurs tests dédiés ; ici, ils s'ouvrent
    assert json.loads((d / "resultat.geojson").read_text(encoding="utf-8"))["type"] == "FeatureCollection"
    assert ET.parse(d / "resultat.gexf").getroot().tag.endswith("gexf")
    from openpyxl import load_workbook

    assert load_workbook(d / "resultat.xlsx").sheetnames == ["Noms", "Détail", "Sources", "Trace"]
    # Webhook : reçu
    assert srv.recus and len(res["envois"]) == 1 and res["envois"][0]["statut"] == 200


class Serveur:
    """Une adresse de test pour le webhook : enregistre ce qu'elle reçoit."""

    recus: ClassVar[list] = []

    def __init__(self, statut=200):
        outer = self
        self.recus = []

        class H(BaseHTTPRequestHandler):
            def do_POST(self):
                outer.recus.append((self.path, self.headers.get("Content-Type"), json.loads(self.rfile.read(int(self.headers["Content-Length"])).decode("utf-8"))))
                self.send_response(statut)
                self.end_headers()

            def log_message(self, *a):
                pass

        self.s = HTTPServer(("127.0.0.1", 0), H)
        threading.Thread(target=self.s.serve_forever, daemon=True).start()
        self.url = f"http://127.0.0.1:{self.s.server_port}/hook"

    def arreter(self):
        self.s.shutdown()


# --- [ ] Le GeoJSON s'ouvre dans uMap ou QGIS et les points tombent au bon endroit.
def test_le_geojson_est_conforme_et_les_points_tombent_au_bon_endroit(monde):
    c, _ = monde
    d = dossier(exporter(c, [{"format": "geojson"}]))
    g = json.loads((d / "resultat.geojson").read_text(encoding="utf-8"))
    # RFC 7946 : FeatureCollection de points [longitude, latitude]
    assert g["type"] == "FeatureCollection" and all(f["type"] == "Feature" and f["geometry"]["type"] == "Point" for f in g["features"])
    attendu = {n[0]: n[5] for n in NOMS if n[5]}
    pts = {f["properties"]["nom"]: f["geometry"]["coordinates"] for f in g["features"]}
    assert set(pts) == {"Fenix Toulouse Handball", "Montpellier Handball", "Handball Club Grenoble"}
    for nom, (lat, lon) in attendu.items():
        assert pts[nom] == [lon, lat]  # longitude d'abord : Toulouse est à 1,44 E et 43,60 N
        assert -180 <= pts[nom][0] <= 180 and -90 <= pts[nom][1] <= 90
    # un point de Toulouse tombe dans un cadre autour de Toulouse, pas au large de l'Afrique (l'erreur classique : latitude et longitude inversées)
    lon, lat = pts["Fenix Toulouse Handball"]
    assert 43.5 < lat < 43.7 and 1.3 < lon < 1.6
    assert g["trace"]["points"] == 3 and g["trace"]["sans_coordonnees"] == 1  # le nom gardé sans coordonnées n'est pas un point, et on le dit


# --- [ ] Le GEXF s'ouvre dans Gephi et affiche les liens.
def test_le_gexf_est_un_graphe_valide_avec_ses_liens(monde):
    c, _ = monde
    d = dossier(exporter(c, [{"format": "gexf"}]))
    racine = ET.parse(d / "resultat.gexf").getroot()
    ns = {"g": "http://gexf.net/1.3"}
    assert racine.get("version") == "1.3" and racine.tag == "{http://gexf.net/1.3}gexf"
    noeuds = {n.get("id"): n.get("label") for n in racine.findall(".//g:nodes/g:node", ns)}
    liens = [(e.get("source"), e.get("target"), e.get("label")) for e in racine.findall(".//g:edges/g:edge", ns)]
    assert len(noeuds) == len(set(noeuds.values())) and "Fenix Toulouse Handball" in noeuds.values()
    assert {noeuds[s]: (noeuds[t], lab) for s, t, lab in liens} == {"Fenix Toulouse Handball": ("Groupe Fenix", "maison mere"), "Handball Club Grenoble": ("Salle des Alpes", "meme adresse")}
    assert all(s in noeuds and t in noeuds for s, t, _ in liens)  # aucun lien ne pointe dans le vide
    ids_edges = [e.get("id") for e in racine.findall(".//g:edges/g:edge", ns)]
    assert len(ids_edges) == len(set(ids_edges))
    declares = {a.get("id") for a in racine.findall(".//g:attributes/g:attribute", ns)}
    assert {v.get("for") for v in racine.findall(".//g:attvalue", ns)} <= declares  # les attributs utilisés sont déclarés
    ext = next(n for n in racine.findall(".//g:nodes/g:node", ns) if n.get("label") == "Groupe Fenix")
    assert any(v.get("value") == "extérieur à la recherche" for v in ext.findall(".//g:attvalue", ns))


# --- [ ] resultat_tsites.csv est accepté par trouver-sites SANS aucune retouche manuelle.
@pytest.mark.reseau
def test_resultat_tsites_est_accepte_par_trouver_sites_sans_retouche(monde, tmp_path):
    c, _ = monde
    if not Path("/opt/trouver-sites/pipeline/tsites/importer.py").exists():
        pytest.skip("trouver-sites absent")
    d = dossier(exporter(c, [{"format": "csv", "contenus": ["tsites"]}]))
    base = tmp_path / "ts"
    base.mkdir()
    sortie = subprocess.run(
        ["python3", "-c", f"from tsites import db; from tsites.importer import importer; cx = db.connexion(); db.migrer(cx); importer(cx, {str(d / 'resultat_tsites.csv')!r})"],
        capture_output=True, text=True, check=False, env={"PATH": "/usr/bin:/bin", "PYTHONPATH": "/opt/trouver-sites/pipeline", "TSITES_DB": str(base / "tsites.db"), "HOME": str(tmp_path)},
    )  # fmt: skip
    assert sortie.returncode == 0, sortie.stdout + sortie.stderr
    assert "occurrences créées   : 4" in sortie.stdout and "lignes lues          : 4" in sortie.stdout
    assert "attention" not in sortie.stdout and "ignorées" not in sortie.stdout and "pays inconnus" not in sortie.stdout  # aucune colonne en trop, aucun pays refusé


# --- [ ] Le webhook envoie bien le résultat à une adresse de test.
def test_le_webhook_envoie_le_resultat_avec_sa_trace(monde):
    c, _ = monde
    srv = Serveur()
    try:
        res = exporter(c, [{"format": "webhook", "contenus": ["detail"], "options": {"url": srv.url}}])
        chemin, type_, corps = srv.recus[0]
    finally:
        srv.arreter()
    assert chemin == "/hook" and type_.startswith("application/json") and res["envois"][0]["noms"] == 4
    verifier_trace(corps["trace"], 4)
    assert {x["nom"] for x in corps["noms"]} == {n[0] for n in NOMS[:3]} | {NOMS[4][0]}
    # l'adresse peut venir de Réglages ; une adresse fausse ou morte est dite, sans empêcher les autres formats
    assert c.put("/api/reglages/sortie", json={"webhook_url": "pas une adresse"}).status_code == 422
    mort = Serveur()
    mort.arreter()
    res2 = exporter(c, [{"format": "webhook", "options": {"url": mort.url}}, {"format": "json"}])
    assert res2["erreurs"] and "ne répond pas" in res2["erreurs"][0]["raison"] and "resultat.json" in res2["fichiers"]
    res3 = exporter(c, [{"format": "webhook", "options": {"url": "ftp://x"}}])
    assert "http://" in res3["erreurs"][0]["raison"]
    res4 = exporter(c, [{"format": "webhook"}])  # aucune adresse ni dans la demande ni dans Réglages
    assert res4["erreurs"] and "http://" in res4["erreurs"][0]["raison"]
    assert c.put("/api/reglages/sortie", json={"webhook_url": srv.url}).status_code == 200
    ok = Serveur()
    c.put("/api/reglages/sortie", json={"webhook_url": ok.url, "actifs": ["webhook"]})
    try:
        r5 = c.post("/api/recherches/r14/export").json()  # sans rien préciser : les formats et l'adresse de Réglages
        assert r5["envois"][0]["adresse"] == ok.url and ok.recus
    finally:
        ok.arreter()


# --- [ ] Ajouter un format est un NOUVEAU FICHIER seulement : le prouver en ajoutant un format bidon sans modifier le cœur.
def test_un_format_est_un_fichier_a_ajouter_sans_toucher_au_coeur(monde):
    c, cfg = monde

    def empreinte():
        return hashlib.sha256(b"".join(f.read_bytes() for f in sorted((RACINE / "src/pilote").rglob("*.py")))).hexdigest()

    avant = empreinte()
    cfg.dossier_formats_perso.mkdir(parents=True, exist_ok=True)
    (cfg.dossier_formats_perso / "bidon.py").write_text(
        'FORMAT = {"id": "bidon", "nom": "Format bidon", "description": "Un format de test.", "extension": ".bidon", "famille": "fichier", "contenus": ["nom"], "contenus_defaut": ["nom"]}\n'
        'def ecrire(jeu, contenu, options, chemin):\n    chemin.write_text("\\n".join("BIDON " + r["nom"] for r in jeu["lignes"]), encoding="utf-8")\n    return {"fichiers": [chemin]}\n', encoding="utf-8")
    (cfg.dossier_formats_perso / "casse.py").write_text('FORMAT = {"id": "casse", "nom": "Cassé", "description": "pas d\'écriture"}\n', encoding="utf-8")  # un format qui ne sait pas écrire n'en est pas un
    ids = {x["id"]: x for x in c.get("/api/formats").json()}
    assert "bidon" in ids and ids["bidon"].get("perso") is True and "casse" not in ids
    res = exporter(c, [{"format": "bidon"}])
    assert (dossier(res) / "resultat.bidon").read_text(encoding="utf-8").splitlines()[0].startswith("BIDON ") and "resultat.bidon" in res["fichiers"]
    assert c.put("/api/reglages/sortie", json={"actifs": ["bidon"]}).status_code == 200  # il est aussi proposé dans Réglages
    assert empreinte() == avant  # pas une ligne du programme n'a changé


# --- [ ] Réexporter une recherche déjà faite dans un autre format ne déclenche AUCUNE requête vers une source.
def test_reexporter_ne_fait_aucune_requete(monde, monkeypatch):
    c, cfg = monde
    import socket
    import urllib.request

    connexions = []
    vrai = socket.socket.connect
    monkeypatch.setattr(socket.socket, "connect", lambda self, adresse, *a, **k: (connexions.append(adresse), vrai(self, adresse, *a, **k))[1])
    monkeypatch.setattr(urllib.request, "urlopen", lambda *a, **k: pytest.fail("une requête réseau part pendant l'export"))
    avant = sorted(p.name for p in cfg.dossier_recherche("r14").glob("*"))
    j_avant = len(c.get("/api/recherches/r14/journal").json())
    for formats_ in ([{"format": "csv"}], [{"format": "geojson"}, {"format": "gexf"}], [{"format": "xlsx"}, {"format": "json"}, {"format": "texte"}]):
        exporter(c, formats_)
        exporter(c, formats_, filtre="tout")
    assert connexions == []  # aucune connexion réseau, vers aucune source
    apres = sorted(p.name for p in cfg.dossier_recherche("r14").glob("*"))
    assert [p for p in apres if p not in avant] == ["export"]  # la recherche n'a rien collecté (ni suivi d'appels, ni journal d'appels) : seulement le dossier de sortie
    journal = c.get("/api/recherches/r14/journal").json()
    assert len(journal) == j_avant + 6 and {x["evenement"] for x in journal[j_avant:]} == {"export"}
    # une recherche rejouée avec un autre filtre donne un autre résultat, depuis la même base
    gardes = json.loads((dossier(exporter(c, [{"format": "json"}], filtre="gardes")) / "resultat.json").read_text(encoding="utf-8"))
    tout = json.loads((dossier(exporter(c, [{"format": "json"}], filtre="tout")) / "resultat.json").read_text(encoding="utf-8"))
    assert len(gardes["noms"]) == 4 and len(tout["noms"]) == 6


# --- [ ] Chaque fichier contient sa date, son sujet et ses sources.
def test_chaque_fichier_porte_sa_trace(monde):
    c, _ = monde
    from openpyxl import load_workbook

    res = exporter(c, [{"format": f} for f in ("csv", "json", "texte", "geojson", "gexf", "xlsx")])
    d = dossier(res)
    # CSV : un fichier voisin (un commentaire gênerait l'outil qui lit le CSV)
    for csv_ in ("resultat.csv", "resultat_tsites.csv"):
        verifier_trace(json.loads((d / f"{csv_}.trace.json").read_text(encoding="utf-8")), 4)
    verifier_trace(json.loads((d / "resultat.json").read_text(encoding="utf-8"))["trace"], 4)
    verifier_trace(json.loads((d / "resultat.geojson").read_text(encoding="utf-8"))["trace"], 4)
    txt = (d / "resultat.txt").read_text(encoding="utf-8")
    assert "Sujet cherché : clubs de handball" in txt and "Recherche : r14" in txt and "Sources utilisées : " in txt and "Exporté le 20" in txt
    desc = ET.parse(d / "resultat.gexf").getroot().find(".//{http://gexf.net/1.3}description").text
    assert "Sujet cherché : clubs de handball" in desc and "Sources utilisées" in desc and "Exporté le 20" in desc
    trace = [r[0] for r in load_workbook(d / "resultat.xlsx")["Trace"].iter_rows(min_row=2, values_only=True)]
    assert any("Sujet cherché : clubs de handball" in x for x in trace) and any("Sources utilisées" in x for x in trace) and any("Exporté le 20" in x for x in trace)


# --- Règles validées : contenu au choix, chemin choisi, décisions, deux fichiers en même temps.
def test_contenu_au_choix_et_deux_fichiers_en_meme_temps(monde):
    c, _ = monde
    res = exporter(c, [{"format": "csv", "contenus": ["nom", "detail"]}])
    d = dossier(res)
    assert next(csv.reader((d / "resultat.csv").read_text(encoding="utf-8").splitlines())) == ["NOM"]  # l'un pour l'outil d'après
    detail = list(csv.reader((d / "resultat_detail.csv").read_text(encoding="utf-8").splitlines(), delimiter=";"))  # l'autre pour l'utilisateur
    assert detail[0] == ["nom", "ville", "pays", "origine", "confirmations", "fonction", "genre", "site", "numero_officiel", "decision", "motif_rejet", "sources", "relations"]
    fenix = next(r for r in detail[1:] if r[0].startswith("Fenix"))
    assert fenix[3] == "dans un registre officiel" and fenix[5] == "siège" and fenix[9] == "gardé" and "Groupe Fenix" in fenix[12]
    assert c.post("/api/recherches/r14/export", json={"formats": [{"format": "csv", "contenus": ["inconnu"]}]}).status_code == 422
    assert c.post("/api/recherches/r14/export", json={"formats": [{"format": "nexiste-pas"}]}).status_code == 422
    assert c.post("/api/recherches/r14/export", json={"filtre": "n'importe quoi"}).status_code == 422


def test_par_defaut_seuls_les_noms_gardes_sortent_et_on_peut_tout_sortir(monde):
    c, _ = monde
    par_defaut = json.loads((dossier(exporter(c, [{"format": "json", "contenus": ["detail"]}])) / "resultat.json").read_text(encoding="utf-8"))
    assert {x["decision"] for x in par_defaut["noms"]} == {"garder"} and par_defaut["trace"]["filtre"] == "noms gardés"
    tout = json.loads((dossier(exporter(c, [{"format": "json"}], filtre="tout")) / "resultat.json").read_text(encoding="utf-8"))
    assert {x["decision"] for x in tout["noms"]} == {"garder", "rejeter", "plus_tard"} and tout["trace"]["filtre"] == "tous les noms, décidés ou non"
    rejetes = json.loads((dossier(exporter(c, [{"format": "json"}], filtre=["rejeter", "plus_tard"])) / "resultat.json").read_text(encoding="utf-8"))
    assert {x["nom"] for x in rejetes["noms"]} == {NOMS[5][0], NOMS[3][0]} and next(x for x in rejetes["noms"] if x["decision"] == "rejeter")["motif_rejet"] == "doublon"
    assert "rejetés" in rejetes["trace"]["filtre"] and "mis de côté" in rejetes["trace"]["filtre"]


def test_le_chemin_est_choisi_par_l_utilisateur_jamais_en_dur(monde, tmp_path):
    c, cfg = monde
    choisi = tmp_path / "mes sorties" / "handball"
    res = exporter(c, [{"format": "json"}], dossier=str(choisi))
    assert dossier(res) == choisi and (choisi / "resultat.json").exists() and (choisi / "decisions.csv").exists()
    # le dossier de Réglages vaut pour toute recherche, relatif au dossier de données s'il n'est pas absolu
    assert c.put("/api/reglages/sortie", json={"dossier": "mes-resultats"}).status_code == 200
    res2 = exporter(c, [{"format": "json"}])
    assert dossier(res2) == cfg.donnees / "mes-resultats" and (cfg.donnees / "mes-resultats" / "resultat.json").exists()
    # sans choix : le dossier de la recherche
    c.put("/api/reglages/sortie", json={"dossier": ""})
    assert dossier(exporter(c, [{"format": "json"}])) == cfg.dossier_recherche("r14") / "export"
    # un dossier impossible à créer est dit en français
    (tmp_path / "fichier").write_text("x", encoding="utf-8")
    r = c.post("/api/recherches/r14/export", json={"formats": [{"format": "json"}], "dossier": str(tmp_path / "fichier" / "sous")})
    assert r.status_code == 422 and "ne peut pas être créé" in r.json()["detail"]


def test_reglages_de_sortie_contenus_et_formats_par_defaut(monde):
    c, _ = monde
    r = c.put("/api/reglages/sortie", json={"actifs": ["csv", "json"], "contenus": {"csv": ["tsites", "detail"]}})
    assert r.status_code == 200 and r.json()["contenus"] == {"csv": ["tsites", "detail"]}
    res = c.post("/api/recherches/r14/export").json()
    assert {"resultat_tsites.csv", "resultat_detail.csv", "resultat.json"} <= set(res["fichiers"]) and "resultat.csv" not in res["fichiers"]
    assert c.put("/api/reglages/sortie", json={"contenus": {"csv": ["n-importe-quoi"]}}).status_code == 422
    assert c.put("/api/reglages/sortie", json={"actifs": ["nexiste-pas"]}).status_code == 422


# --- Sécurité : une cellule qui commencerait par = serait une formule dans un tableur.
def test_une_formule_dans_un_nom_est_neutralisee(monde):
    c, _ = monde
    from openpyxl import load_workbook

    d = dossier(exporter(c, [{"format": "csv", "contenus": ["detail"]}, {"format": "xlsx"}]))
    noms = [r[0] for r in csv.reader((d / "resultat_detail.csv").read_text(encoding="utf-8").splitlines(), delimiter=";")][1:]
    assert any(n.startswith("'=HYPERLINK") for n in noms) and not any(n.startswith("=") for n in noms)
    ws = load_workbook(d / "resultat.xlsx")["Noms"]
    cellules = [c_ for ligne in ws.iter_rows(min_row=2) for c_ in ligne if isinstance(c_.value, str) and "HYPERLINK" in c_.value]
    assert cellules and all(c_.data_type == "s" and not c_.value.startswith("=") for c_ in cellules)


def test_un_caractere_de_controle_dans_un_nom_ne_casse_pas_le_gexf(cfg):
    fabrique = creer_fabrique(cfg.url_base)
    with fabrique.begin() as s:
        s.add(Recherche(id="rc", qui="sednal", sujet={"quoi": "x"}, parametres={}, etat="terminee"))
        s.flush()
        s.add(Candidat(recherche_id="rc", nom="Club\x07 Étrange <&> \"A\"", genre_source="web", trouve_par=["searxng"], nb_endroits=1, faits={}, relations=[{"type": "lie", "vers": "Autre\x0b nom"}]))
    with TestClient(creer_app(cfg)) as c:
        ids = [x["id"] for x in c.get("/api/recherches/rc/tri").json()["lignes"]]
        c.post("/api/recherches/rc/decisions", json={"valeur": "garder", "ids": ids})
        res = c.post("/api/recherches/rc/export", json={"formats": [{"format": "gexf"}, {"format": "xlsx"}]}).json()
    labels = [n.get("label") for n in ET.parse(Path(res["dossier"]) / "resultat.gexf").getroot().iter("{http://gexf.net/1.3}node")]
    assert labels == ['Club Étrange <&> "A"', "Autre nom"]  # le XML reste valide, les signes spéciaux sont échappés, le reste du nom est intact


def test_les_fichiers_d_une_sortie_n_ont_pas_de_bom_ni_de_retour_chariot(monde):
    c, _ = monde
    d = dossier(exporter(c, [{"format": "csv", "contenus": ["tsites"]}]))
    brut = (d / "resultat_tsites.csv").read_bytes()
    assert not brut.startswith(b"\xef\xbb\xbf") and b"\r" not in brut and brut.endswith(b"\n")
    assert io.StringIO(brut.decode("utf-8")).readline().strip() == "nom;ville;pays"


# --- L'export se fait depuis la page Tri (vrai navigateur), et les choix se règlent dans Réglages.
@pytest.fixture
def page_tri(chromium, serveur_vivant, cfg):
    peupler(cfg)
    with TestClient(creer_app(cfg)) as c:
        c.post("/api/assistant/terminer", json={"etape1": {"pays": "FR", "presets_coches": ["web"], "chercher": [], "plafond": 5}, "etape2": {"active": True, "mode": "memes", "verifier": []}, "etape3": {"actifs": ["csv"], "dossier": ""}})
        decider(c)
    url = serveur_vivant(creer_app(cfg))
    ctx = chromium.new_context(viewport={"width": 1400, "height": 1000})
    page = ctx.new_page()
    erreurs = []
    page.on("pageerror", lambda e: erreurs.append(str(e)))
    page.goto(url + "/")
    page.evaluate("localStorage.setItem('g_tri_recherche','r14')")
    page.reload()
    yield page, cfg
    ctx.close()
    assert erreurs == [], erreurs


def test_exporter_depuis_la_page_tri_dans_plusieurs_formats(page_tri, tmp_path):
    page, _ = page_tri
    page.click('.tab[data-p="tri"]')
    page.wait_for_selector("#export-formats input[data-f]")
    assert page.locator("#export-formats input[data-f]").count() >= 7  # chaque format du dossier formats/ a sa case : rien n'est écrit en dur dans la page
    assert page.locator('#export-formats input[data-f="csv"]').is_checked()  # ceux de Réglages sont cochés
    page.check('#export-formats input[data-f="json"]')
    page.check('#export-formats input[data-f="geojson"]')
    page.check('#export-formats input[data-fc="csv"][data-c="detail"]')
    cible = tmp_path / "sorties-test"
    page.fill("#export-dossier", str(cible))
    page.click("#exporter")
    page.wait_for_function("document.querySelector('#export-msg').textContent.includes('exporté')", timeout=20000)
    msg = page.inner_text("#export-msg")
    assert "4 noms exportés (noms gardés)" in msg and "resultat_detail.csv" in msg and "resultat.geojson" in msg and str(cible) in msg
    assert (cible / "resultat_detail.csv").exists() and (cible / "resultat.json").exists() and (cible / "resultat.geojson").exists()
    # un autre choix de noms, sans rien relancer
    page.select_option("#export-filtre", "tout")
    page.click("#exporter")
    page.wait_for_function("document.querySelector('#export-msg').textContent.includes('6 noms exportés')", timeout=20000)
    assert len(json.loads((cible / "resultat.json").read_text(encoding="utf-8"))["trace"]["sources"]) >= 3
    # le webhook demande son adresse
    page.check('#export-formats input[data-f="webhook"]')
    assert page.locator("#export-webhook-w").is_visible()


def test_reglages_de_sortie_dans_la_page_reglages(page_tri):
    page, cfg = page_tri
    page.click('.tab[data-p="reglages"]')
    page.evaluate("document.querySelectorAll('details.bloc').forEach(d => d.open = true)")  # les blocs de Réglages sont repliés au départ
    page.wait_for_selector("#list3 .srow")
    assert page.locator("#list3 input[data-id]").count() >= 7 and page.locator('#list3 input[data-fc="csv"]').count() == 3  # trois contenus pour le CSV
    page.uncheck('#list3 input[data-fc="csv"][data-c="tsites"]')
    import yaml

    for _ in range(100):  # le réglage part dès que la case est décochée : on attend qu'il soit écrit
        if cfg.fichier_reglages.exists() and (yaml.safe_load(cfg.fichier_reglages.read_text(encoding="utf-8")) or {}).get("sortie", {}).get("contenus"):
            break
        page.wait_for_timeout(100)
    page.fill("#sortie-dossier", "mes-sorties")
    page.fill("#sortie-webhook", "https://exemple.example/hook")
    page.click("#sortie-ok")
    page.wait_for_function("document.querySelector('#sortie-msg').textContent.includes('Enregistré')")
    y = yaml.safe_load(cfg.fichier_reglages.read_text(encoding="utf-8"))["sortie"]
    assert y["contenus"]["csv"] == ["nom"] and y["dossier"] == "mes-sorties" and y["webhook_url"] == "https://exemple.example/hook" and y["actifs"] == ["csv"]
    page.click('#list3 label.sw:has(input[data-id="json"])')  # activer un format ne touche pas aux contenus
    for _ in range(100):
        if "json" in yaml.safe_load(cfg.fichier_reglages.read_text(encoding="utf-8"))["sortie"]["actifs"]:
            break
        page.wait_for_timeout(100)
    y2 = yaml.safe_load(cfg.fichier_reglages.read_text(encoding="utf-8"))["sortie"]
    assert y2["actifs"] == ["csv", "json"] and y2["contenus"]["csv"] == ["nom"] and "None" not in str(y2)


def test_la_sortie_choisie_dans_recherches_est_appliquee_a_l_export_de_cette_recherche(cfg, tmp_path):
    peupler(cfg)
    fabrique = creer_fabrique(cfg.url_base)
    with fabrique.begin() as s:
        r = s.get(Recherche, "r14")
        r.parametres = {**r.parametres, "sortie": {"actifs": ["texte"], "dossier": str(tmp_path / "pour-celle-ci")}}
    with TestClient(creer_app(cfg)) as c:
        decider(c)
        res = c.post("/api/recherches/r14/export").json()  # rien de précisé : les choix de cette recherche, pas ceux de Réglages (csv)
    assert dossier(res) == tmp_path / "pour-celle-ci" and (tmp_path / "pour-celle-ci" / "resultat.txt").exists() and not (tmp_path / "pour-celle-ci" / "resultat.csv").exists()
